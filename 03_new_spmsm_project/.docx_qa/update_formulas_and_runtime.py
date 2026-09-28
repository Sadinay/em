from pathlib import Path
from copy import deepcopy

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


DOCX = Path(r"C:\Users\26096\Desktop\SPMSM_论文大纲与研究进度_修订版.docx")
QA = Path(r"C:\Users\26096\Desktop\em\03_new_spmsm_project\.docx_qa")
EQ_DIR = QA / "equations_100gen"
BACKUP = QA / "outline_before_formula_runtime_update.docx"
EQ_DIR.mkdir(parents=True, exist_ok=True)


FORMULAS = {
    "eq_6_1_metrics": r"T_{\mathrm{avg}}(\mathbf{x})=\frac{1}{6}\sum_{k=1}^{6}T_k(\mathbf{x}),\qquad \Delta T(\mathbf{x})=\max_{1\leq k\leq6}T_k(\mathbf{x})-\min_{1\leq k\leq6}T_k(\mathbf{x})",
    "eq_6_1_objective": r"\min_{\mathbf{x}\in\{0,1\}^{120}}\left(f_1(\mathbf{x}),f_2(\mathbf{x})\right)=\left(-T_{\mathrm{avg}}(\mathbf{x}),\Delta T(\mathbf{x})\right)",
    "eq_6_1_constraints": r"12\leq\sum_{i=1}^{120}x_i\leq108,\quad x_i\in\{0,1\},\qquad p_c=0.9,\quad p_m=\frac{1}{120}",
    "eq_6_6_mean": r"\mu_t(\mathbf{x})=\frac{1}{4}\sum_{m=1}^{4}\widehat{y}_{m,t}(\mathbf{x}),\qquad t\in\{T_{\mathrm{avg}},\Delta T\}",
    "eq_6_6_std": r"d_t(\mathbf{x})=\sqrt{\frac{1}{4}\sum_{m=1}^{4}\left[\widehat{y}_{m,t}(\mathbf{x})-\mu_t(\mathbf{x})\right]^2}",
    "eq_6_6_combined": r"D(\mathbf{x})=\sqrt{\frac{1}{2}\left[\left(\frac{d_{T_{\mathrm{avg}}}(\mathbf{x})}{s_{T_{\mathrm{avg}}}}\right)^2+\left(\frac{d_{\Delta T}(\mathbf{x})}{s_{\Delta T}}\right)^2\right]}",
    "eq_7_3_improvement": r"r_T=\frac{4.2593-4.1082}{4.1082}\times100\%=3.68\%,\qquad r_{\Delta T}=\frac{0.0802-0.0076}{0.0802}\times100\%=90.52\%",
    "eq_7_6_counts": r"N_{\mathrm{full}}=5540\times6=33240,\quad N_{\mathrm{actual}}=601\times6=3606,\quad \eta_N=1-\frac{601}{5540}=89.15\%",
    "eq_7_6_time": r"t_{\mathrm{full}}\approx10.74\,\mathrm{h},\quad t_{\mathrm{actual}}=1.58\,\mathrm{h},\quad \Delta t\approx9.16\,\mathrm{h}\;(85.3\%)",
}


def render_formula(name: str, formula: str) -> Path:
    out = EQ_DIR / f"{name}.png"
    width = 12.0
    fig = plt.figure(figsize=(width, 0.72), facecolor="white")
    fig.text(0.5, 0.5, f"${formula}$", ha="center", va="center", fontsize=17, color="black")
    plt.axis("off")
    fig.savefig(out, dpi=360, bbox_inches="tight", pad_inches=0.06, facecolor="white")
    plt.close(fig)
    return out


def paragraph_after(paragraph):
    new_p = deepcopy(paragraph._p)
    for child in list(new_p):
        new_p.remove(child)
    paragraph._p.addnext(new_p)
    return paragraph.__class__(new_p, paragraph._parent)


def set_text(paragraph, text: str):
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.name = "宋体"
    run.font.size = Pt(10.5)


def add_equation_after(paragraph, image_path: Path, width: float = 6.2):
    p = paragraph_after(paragraph)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(image_path), width=Inches(width))
    return p


images = {name: render_formula(name, formula) for name, formula in FORMULAS.items()}

if not BACKUP.exists():
    BACKUP.write_bytes(DOCX.read_bytes())

doc = Document(DOCX)

def find_exact(prefix: str):
    for p in doc.paragraphs:
        if p.text.startswith(prefix):
            return p
    raise RuntimeError(f"Paragraph not found: {prefix}")


p = find_exact("设计变量为6×20排列")
set_text(p, "设计变量为6×20排列的120位二值材料基因（0为空气，1为磁体）。代理与FEMM采用同一工况下的平均转矩与六角度绝对转矩波动，单位均为N·m，定义为：")
add_equation_after(p, images["eq_6_1_metrics"], 6.25)

p = find_exact("首轮采用二值 NSGA-II")
set_text(p, "首轮采用二值NSGA-II，将平均转矩最大化和绝对转矩波动最小化作为两个独立目标：")
eqp = add_equation_after(p, images["eq_6_1_objective"], 5.8)
note = paragraph_after(eqp)
set_text(note, "当前磁体格数12—108仅作为搜索边界，来源于既有候选生成范围，不等同于已验证的制造约束。遗传算子参数与搜索边界统一写为：")
add_equation_after(note, images["eq_6_1_constraints"], 5.7)

p = find_exact("首轮正式运行采用种群96")
set_text(p, "首轮正式运行采用种群96、100代和随机种子20260921；每代最多新增10个FEMM验证基因。Polar90 ResNet18负责主性能预测，四个保留模型共同提供分歧信息。")

p = find_exact("μ_t(x)")
p.clear()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(images["eq_6_6_mean"]), width=Inches(5.5))
p.paragraph_format.space_after = Pt(4)

p = find_exact("d_t(x)")
p.clear()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(images["eq_6_6_std"]), width=Inches(5.15))
p.paragraph_format.space_after = Pt(4)

p = find_exact("这里把四个已选模型视为固定模型组")
set_text(p, "这里把四个已选模型视为固定模型组，因此分母为4，不沿用样本标准差的M−1。两个目标先分别计算分歧；需要综合选点时，再按训练集目标标准差归一化：")
eqp = add_equation_after(p, images["eq_6_6_combined"], 5.7)
note = paragraph_after(eqp)
set_text(note, "各模型在冻结验证集上的误差另行记录，不预先混入分歧公式。D(x)仅是四模型分歧的启发式指标，不是已校准的误差界。")

p = find_exact("D(x) 是四模型分歧")
set_text(p, "选点同时覆盖预测帕累托前沿两端与中部、前沿附近的高分歧结构，以及少量结构覆盖和随机审计样本。首轮在第1、20、40、60、80和100代保存FEMM真实前沿，用于观察搜索推进。")

p = find_exact("第1代至第100代")
set_text(p, "第1代至第100代，最大平均转矩由4.1082升至4.2593 N·m，最低转矩波动由0.0802降至0.0076 N·m；第60代后转矩上限基本稳定，后续改进主要来自低波动区域。对应的转矩提升率与波动下降率为：")
add_equation_after(p, images["eq_7_3_improvement"], 5.9)

p = find_exact("100代共验证601个不同基因")
set_text(p, "100代共生成5,540个去重候选，其中601个完成六角度FEMM验证。已验证基因直接复用真值，实际求解量及相对全FEMM方案的降幅为：")
eqp = add_equation_after(p, images["eq_7_6_counts"], 6.05)
note = paragraph_after(eqp)
set_text(note, "运行日志显示，一百代端到端实际用时约1小时35分。按同一硬件、六角度设置、5个worker理想并行，并保持约0.47小时固定开销不变，若5,540个候选全部使用FEMM，预计需约10小时44分：")
eq2 = add_equation_after(note, images["eq_7_6_time"], 5.45)
note2 = paragraph_after(eq2)
set_text(note2, "因此，FEMM求解数量确定减少89.15%；端到端预计节省约9小时09分，时间减少约85.3%。时间结果属于基于本机日志的工程估算，会随硬件、并行效率和I/O竞争变化。")

doc.save(DOCX)
print(DOCX)
